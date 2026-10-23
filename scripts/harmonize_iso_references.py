# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def harmonize_iso_references():
    doc_path = 'docs/gestion_academica/entregas/INFORME_ISO25010_SEMANAS_10_11_LEOFIT.docx'
    doc = docx.Document(doc_path)

    # Buscar el encabezado de "Referencias y marco de consulta"
    ref_idx = None
    for i, p in enumerate(doc.paragraphs):
        if 'referencias y marco de consulta' in p.text.lower():
            ref_idx = i
            break

    if ref_idx is None:
        print("No se encontró la sección de referencias.")
        return

    print(f"Encabezado de referencias encontrado en párrafo {ref_idx}: {doc.paragraphs[ref_idx].text}")

    # Encontrar el siguiente encabezado (ANEXO A)
    next_section_idx = None
    for i in range(ref_idx + 1, len(doc.paragraphs)):
        if doc.paragraphs[i].text.strip().startswith('ANEXO A'):
            next_section_idx = i
            break

    print(f"Siguiente sección (ANEXO A) en párrafo {next_section_idx}")

    # Párrafos a eliminar entre ref_idx y next_section_idx
    paras_to_remove = doc.paragraphs[ref_idx + 1:next_section_idx]
    for p in paras_to_remove:
        p._element.getparent().remove(p._element)

    # Nuevas referencias completas bajo APA 7ma con URLs canónicas verificadas
    # Estructura: (autor_anio, titulo_cursiva, resto, url)
    apa_refs = [
        (
            "IEEE Computer Society. (1998). ",
            "IEEE recommended practice for software requirements specifications",
            " (IEEE Std 830-1998). Institute of Electrical and Electronics Engineers. ",
            "https://doi.org/10.1109/IEEESTD.1998.88286"
        ),
        (
            "ISO/IEC. (2016). ",
            "Systems and software engineering — Systems and software Quality Requirements and Evaluation (SQuaRE) — Measurement of system and software product quality",
            " (ISO/IEC 25023:2016). International Organization for Standardization. ",
            "https://www.iso.org/standard/35747.html"
        ),
        (
            "ISO/IEC. (2023). ",
            "Systems and software engineering — Systems and software quality requirements and evaluation (SQuaRE) — Product quality model",
            " (ISO/IEC 25010:2023). International Organization for Standardization. ",
            "https://www.iso.org/standard/78176.html"
        ),
        (
            "OWASP Foundation. (2021). ",
            "OWASP Top 10:2021 — The ten most critical web application security risks",
            ". Open Web Application Security Project. ",
            "https://owasp.org/Top10/"
        ),
        (
            "PostgreSQL Global Development Group. (2024). ",
            "PostgreSQL 16 documentation",
            ". The PostgreSQL Global Development Group. ",
            "https://www.postgresql.org/docs/16/"
        ),
        (
            "Universidad Tecnológica del Perú. (2026). ",
            "Sílabo del Curso integrador II: software (100000S12F), ciclo 2026-2",
            ". Dirección Académica UTP.",
            None
        ),
        (
            "World Wide Web Consortium [W3C]. (2018). ",
            "Web content accessibility guidelines (WCAG) 2.1",
            ". W3C Recommendation. ",
            "https://www.w3.org/TR/WCAG21/"
        )
    ]

    # Insertar párrafos justo después del encabezado
    ref_heading = doc.paragraphs[ref_idx]
    last_elem = ref_heading._element

    for autor, titulo, resto, url in apa_refs:
        p = doc.add_paragraph()
        doc._body._body.remove(p._element) # Desasociar del final
        last_elem.addnext(p._element) # Insertar en orden
        last_elem = p._element

        p.paragraph_format.left_indent = Inches(0.5) # 1.27 cm
        p.paragraph_format.first_line_indent = Inches(-0.5) # Sangría francesa
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.space_before = Pt(2)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        # Run 1: Autor y fecha
        r1 = p.add_run(autor)
        r1.font.name = 'Arial'
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(30, 41, 59)

        # Run 2: Título en cursiva
        r2 = p.add_run(titulo)
        r2.italic = True
        r2.font.name = 'Arial'
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = RGBColor(30, 41, 59)

        # Run 3: Resto
        r3 = p.add_run(resto)
        r3.font.name = 'Arial'
        r3.font.size = Pt(9.5)
        r3.font.color.rgb = RGBColor(30, 41, 59)

        # Run 4: URL
        if url:
            r4 = p.add_run(url)
            r4.font.name = 'Arial'
            r4.font.size = Pt(9)
            r4.font.color.rgb = RGBColor(30, 58, 138)
            r4.underline = True

    # Agregar párrafo en blanco separador antes de ANEXO A
    p_sep = doc.add_paragraph()
    doc._body._body.remove(p_sep._element)
    last_elem.addnext(p_sep._element)
    p_sep.paragraph_format.space_after = Pt(12)

    doc.save(doc_path)
    print(f"¡Referencias APA 7ma armonizadas exitosamente en {doc_path}!")

if __name__ == '__main__':
    harmonize_iso_references()
