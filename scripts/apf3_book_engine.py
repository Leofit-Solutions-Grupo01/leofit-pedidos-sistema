"""
Motor de maquetación del Informe Final APF3 (formato libro de ingeniería de software).

Cumplimiento normativo estricto:
- Tipografía: Times New Roman, 12 pt para el cuerpo del texto.
- Niveles de títulos claramente diferenciados en Times New Roman.
- Texto y tablas en blanco y negro (escala de grises profesional y formal).
- Figuras, diagramas, capturas y gráficos a todo color en alta resolución.
- Sangría de primera línea de 1.27 cm (0.5 pulgadas) en párrafos de texto continuo.
- Numeración de páginas formal (romanos en front-matter, arábigos desde el inicio del cuerpo).
- Título formal del documento centrado y en negrita en la primera página del texto.
"""
import os
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Inches, RGBColor

# Paleta formal en Blanco y Negro para texto y tablas
BLACK = RGBColor(0x00, 0x00, 0x00)
DARK_GREY = RGBColor(0x26, 0x26, 0x26)
MID_GREY = RGBColor(0x55, 0x55, 0x55)
LIGHT_GREY = RGBColor(0x88, 0x88, 0x88)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Tipografías requeridas
BODY_FONT = "Times New Roman"
HEAD_FONT = "Times New Roman"
CODE_FONT = "Consolas"


def _rfonts(style_or_run_element, name):
    rpr = style_or_run_element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


def shade(cell_or_par_pr, hex_fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    cell_or_par_pr.append(shd)


def cell_bg(cell, hex_fill):
    shade(cell._tc.get_or_add_tcPr(), hex_fill)


def cell_margins(table, top=50, bottom=50, left=90, right=90):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblPr.append(m)


def table_borders(table, color="404040", sz=4):
    """Bordes formales en escala de grises / negro para tablas."""
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trPr.append(e)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit")
    e.set(qn("w:val"), "true")
    trPr.append(e)


def add_field(paragraph, instr, placeholder="", bold=None, size=None, color=None):
    """Campo complejo de Word (begin / instr / separate / resultado / end)."""
    def mk(kind):
        r = paragraph.add_run()
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), kind)
        r._r.append(fc)
        return r
    mk("begin")
    r = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    mk("separate")
    res = paragraph.add_run(placeholder)
    for run in (res,):
        if bold is not None:
            run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
    mk("end")
    return res


def para_border(paragraph, side="bottom", color="000000", sz=8, space=4):
    """Línea de borde formal para encabezados o títulos."""
    pPr = paragraph._p.get_or_add_pPr()
    pb = pPr.find(qn("w:pBdr"))
    if pb is None:
        pb = OxmlElement("w:pBdr")
        pPr.append(pb)
    e = OxmlElement(f"w:{side}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), str(space))
    e.set(qn("w:color"), color)
    pb.append(e)


class Book:
    def __init__(self, root):
        self.root = root
        self.doc = Document()
        self.fig_n = 0
        self.tab_n = 0
        self.chap_n = 0
        self.sec_n = 0
        self._setup_page()
        self._setup_styles()

    # ------------------------------------------------------------------ setup
    def _setup_page(self):
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.54) # 1 pulgada formal
        s.top_margin = s.bottom_margin = Cm(2.54) # 1 pulgada formal
        s.header_distance = Cm(1.27)
        s.footer_distance = Cm(1.27)

    def _setup_styles(self):
        st = self.doc.styles
        n = st["Normal"]
        n.font.name = BODY_FONT
        n.font.size = Pt(12)
        n.font.color.rgb = BLACK
        _rfonts(n.element, BODY_FONT)
        n.paragraph_format.space_after = Pt(6)
        n.paragraph_format.line_spacing = 1.3

        # Configuración formal de niveles de títulos en Times New Roman (Blanco y Negro)
        # Nivel 1: 16 pt, Negrita
        # Nivel 2: 14 pt, Negrita
        # Nivel 3: 13 pt, Negrita
        # Nivel 4: 12 pt, Negrita y Cursiva
        cfg = {
            "Heading 1": (16, BLACK, 18, 10, True, False),
            "Heading 2": (14, BLACK, 14, 8, True, False),
            "Heading 3": (13, BLACK, 12, 6, True, False),
            "Heading 4": (12, BLACK, 8, 4, True, True),
        }
        for name, (sz, col, before, after, is_bold, is_italic) in cfg.items():
            h = st[name]
            h.font.name = HEAD_FONT
            h.font.size = Pt(sz)
            h.font.bold = is_bold
            h.font.italic = is_italic
            h.font.color.rgb = col
            _rfonts(h.element, HEAD_FONT)
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.line_spacing = 1.15

        cap = st["Caption"]
        cap.font.name = HEAD_FONT
        cap.font.size = Pt(10)
        cap.font.bold = False
        cap.font.italic = False
        cap.font.color.rgb = BLACK
        _rfonts(cap.element, HEAD_FONT)
        cap.paragraph_format.space_after = Pt(8)

        for name in ("List Bullet", "List Number"):
            s_ = st[name]
            s_.font.name = BODY_FONT
            s_.font.size = Pt(12)
            s_.font.color.rgb = BLACK
            _rfonts(s_.element, BODY_FONT)
            s_.paragraph_format.space_after = Pt(4)
            s_.paragraph_format.line_spacing = 1.25

        for name in ("TOC 1", "TOC 2", "TOC 3", "Table of Figures"):
            try:
                t = st[name]
            except KeyError:
                from docx.enum.style import WD_STYLE_TYPE
                t = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
                t.base_style = st["Normal"]
            t.font.name = HEAD_FONT
            _rfonts(t.element, HEAD_FONT)
            t.paragraph_format.space_after = Pt(3)
            t.paragraph_format.line_spacing = 1.15

        st["TOC 1"].font.bold = True
        st["TOC 1"].font.size = Pt(12)
        st["TOC 1"].font.color.rgb = BLACK
        st["TOC 1"].paragraph_format.space_before = Pt(8)

        st["TOC 2"].font.size = Pt(11)
        st["TOC 2"].font.color.rgb = BLACK
        st["TOC 2"].paragraph_format.left_indent = Cm(0.5)

        st["TOC 3"].font.size = Pt(10)
        st["TOC 3"].font.color.rgb = BLACK
        st["TOC 3"].paragraph_format.left_indent = Cm(1.2)

        st["Table of Figures"].font.size = Pt(10.5)
        st["Table of Figures"].font.color.rgb = BLACK

    # ------------------------------------------------------------------ texto
    def _inline(self, p, text, size=None, color=None, bold_all=False):
        parts = re.split(r"(\*\*.+?\*\*|`.+?`|_\_.+?_\_)", text)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
                r.font.name = BODY_FONT
                _rfonts(r._r, BODY_FONT)
            elif part.startswith("`") and part.endswith("`"):
                r = p.add_run(part[1:-1])
                r.font.name = CODE_FONT
                _rfonts(r._r, CODE_FONT)
                r.font.size = Pt((size or 12) - 1.5)
                r.font.color.rgb = DARK_GREY
                continue
            else:
                r = p.add_run(part)
                r.font.name = BODY_FONT
                _rfonts(r._r, BODY_FONT)
                if bold_all:
                    r.bold = True
            if size:
                r.font.size = Pt(size)
            r.font.color.rgb = color if color is not None else BLACK
        return p

    def p(self, text, align="justify", size=None, italic=False, after=None, indent=None, first_indent=None):
        """Párrafo de texto con sangría de primera línea de 1.27 cm (0.5 in) según norma formal."""
        par = self.doc.add_paragraph()
        actual_size = size or 12
        self._inline(par, text, size=actual_size, color=BLACK)
        par.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
                         "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
        if italic:
            for r in par.runs:
                r.italic = True
        if after is not None:
            par.paragraph_format.space_after = Pt(after)
        else:
            par.paragraph_format.space_after = Pt(6)
        par.paragraph_format.line_spacing = 1.3

        if indent:
            par.paragraph_format.left_indent = Cm(indent)

        # Sangría de primera línea: 1.27 cm (0.5 pulgadas)
        # Se aplica automáticamente a párrafos de texto continuo justificados
        if first_indent is True or (first_indent is None and align == "justify" and indent is None):
            par.paragraph_format.first_line_indent = Inches(0.5)
        else:
            par.paragraph_format.first_line_indent = Inches(0)

        return par

    def bullets(self, items):
        for it in items:
            par = self.doc.add_paragraph(style="List Bullet")
            par.paragraph_format.first_line_indent = Inches(0)
            self._inline(par, it, size=12, color=BLACK)
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(6)

    def numbered(self, items):
        for i, it in enumerate(items, 1):
            par = self.doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(1.27)
            par.paragraph_format.first_line_indent = Cm(-0.6)
            par.paragraph_format.space_after = Pt(3)
            r = par.add_run(f"{i}.  ")
            r.bold = True
            r.font.name = BODY_FONT
            _rfonts(r._r, BODY_FONT)
            r.font.color.rgb = BLACK
            self._inline(par, it, size=12, color=BLACK)
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(6)

    # ------------------------------------------------------------------ títulos
    def doc_title(self, title):
        """Título formal del documento en la primera página del texto (centrado y en negrita)."""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(20)
        p.paragraph_format.first_line_indent = Inches(0)
        r = p.add_run(title)
        r.bold = True
        r.font.name = HEAD_FONT
        _rfonts(r._r, HEAD_FONT)
        r.font.size = Pt(14)
        r.font.color.rgb = BLACK
        return p

    def part(self, title, intro=None):
        h = self.doc.add_heading(title, level=1)
        h.paragraph_format.page_break_before = True
        h.paragraph_format.first_line_indent = Inches(0)
        h.paragraph_format.space_before = Pt(40)
        para_border(h, "bottom", "000000", 12, 6)
        if intro:
            par = self.p(intro, italic=True, first_indent=False)
            par.paragraph_format.space_before = Pt(12)
        return h

    def front(self, title, page_break=True):
        h = self.doc.add_heading(title, level=1)
        h.paragraph_format.page_break_before = page_break
        h.paragraph_format.first_line_indent = Inches(0)
        h.paragraph_format.space_before = Pt(0)
        para_border(h, "bottom", "000000", 12, 6)
        return h

    def chapter(self, title):
        self.chap_n += 1
        self.sec_n = 0
        h = self.doc.add_heading(f"{self.chap_n}. {title}", level=2)
        h.paragraph_format.page_break_before = True
        h.paragraph_format.first_line_indent = Inches(0)
        h.paragraph_format.space_before = Pt(0)
        para_border(h, "bottom", "333333", 8, 6)
        return h

    def annex(self, letter, title):
        h = self.doc.add_heading(f"Anexo {letter}. {title}", level=2)
        h.paragraph_format.page_break_before = True
        h.paragraph_format.first_line_indent = Inches(0)
        h.paragraph_format.space_before = Pt(0)
        para_border(h, "bottom", "333333", 8, 6)
        self.sec_n = 0
        self.annex_letter = letter
        return h

    def sec(self, title, annex=False):
        self.sec_n += 1
        label = f"{getattr(self, 'annex_letter', '')}.{self.sec_n}" if annex else f"{self.chap_n}.{self.sec_n}"
        h = self.doc.add_heading(f"{label}  {title}", level=3)
        h.paragraph_format.first_line_indent = Inches(0)
        return h

    def sub(self, title):
        h = self.doc.add_heading(title, level=4)
        h.paragraph_format.first_line_indent = Inches(0)
        return h

    # ------------------------------------------------------------------ captions
    def _caption(self, label, n, text, before=False):
        cp = self.doc.add_paragraph(style="Caption")
        cp.paragraph_format.first_line_indent = Inches(0)
        r = cp.add_run(f"{label} ")
        r.bold = True
        r.font.name = HEAD_FONT
        _rfonts(r._r, HEAD_FONT)
        r.font.size = Pt(10)
        r.font.color.rgb = BLACK
        res = add_field(cp, f"SEQ {label} \\* ARABIC", str(n), bold=True, color=BLACK)
        for rx in cp.runs:
            rx.font.name = HEAD_FONT
            _rfonts(rx._r, HEAD_FONT)
        r2 = cp.add_run(f". {text}")
        r2.font.name = HEAD_FONT
        _rfonts(r2._r, HEAD_FONT)
        r2.font.size = Pt(10)
        r2.font.color.rgb = BLACK
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if before:
            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            cp.paragraph_format.keep_with_next = True
            cp.paragraph_format.space_before = Pt(8)
            cp.paragraph_format.space_after = Pt(4)
        return cp

    # ------------------------------------------------------------------ figuras (A COLOR)
    def fig(self, rel_path, caption, width_cm=15.0):
        path = rel_path if os.path.isabs(rel_path) else os.path.join(self.root, rel_path)
        if not os.path.exists(path):
            raise FileNotFoundError(path)
        self.fig_n += 1
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.first_line_indent = Inches(0)
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.space_before = Pt(8)
        par.paragraph_format.space_after = Pt(4)
        par.add_run().add_picture(path, width=Cm(width_cm))
        self._caption("Figura", self.fig_n, caption)

    # ------------------------------------------------------------------ tablas (EN BLANCO Y NEGRO)
    def table(self, headers, rows, caption, widths=None, size=9.5, align=None):
        self.tab_n += 1
        self._caption("Tabla", self.tab_n, caption, before=True)
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        table_borders(t, color="404040", sz=4) # Bordes negros/grises
        cell_margins(t, top=60, bottom=60, left=90, right=90)
        total = 16.0
        if widths is None:
            widths = [total / len(headers)] * len(headers)
        else:
            k = total / sum(widths)
            widths = [w * k for w in widths]

        # Encabezado de tabla en blanco y negro (fondo gris oscuro carbón, texto blanco en negrita)
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.width = Cm(widths[i])
            cell_bg(c, "262626")
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            c.paragraphs[0].paragraph_format.first_line_indent = Inches(0)
            r = c.paragraphs[0].add_run(h)
            r.bold = True
            r.font.size = Pt(size)
            r.font.name = HEAD_FONT
            _rfonts(r._r, HEAD_FONT)
            r.font.color.rgb = WHITE
        repeat_header(t.rows[0])

        # Filas de datos en escala de grises
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            no_split(t.rows[-1])
            for i, val in enumerate(row):
                c = cells[i]
                c.width = Cm(widths[i])
                if ri % 2 == 1:
                    cell_bg(c, "F5F5F5") # Gris muy claro alterno
                par = c.paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.15
                par.paragraph_format.first_line_indent = Inches(0)
                if align and align[i] == "c":
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._inline(par, str(val), size=size, color=BLACK)
                for r in par.runs:
                    if r.font.name != CODE_FONT:
                        r.font.name = HEAD_FONT
                        _rfonts(r._r, HEAD_FONT)
                    if r.font.color.rgb is None or r.font.color.rgb != WHITE:
                        r.font.color.rgb = BLACK
        sp = self.doc.add_paragraph()
        sp.paragraph_format.first_line_indent = Inches(0)
        sp.paragraph_format.space_after = Pt(6)
        return t

    # ------------------------------------------------------------------ código (escala de grises)
    def code(self, text, caption=None, size=8.5):
        if caption:
            self.tab_n += 1
            self._caption("Tabla", self.tab_n, caption, before=True)
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_borders(t, "737373", 4)
        cell_margins(t, 80, 80, 140, 140)
        c = t.rows[0].cells[0]
        c.width = Cm(16)
        cell_bg(c, "F5F5F5")
        lines = text.strip("\n").split("\n")
        for i, ln in enumerate(lines):
            par = c.paragraphs[0] if i == 0 else c.add_paragraph()
            par.paragraph_format.first_line_indent = Inches(0)
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            r = par.add_run(ln if ln else " ")
            r.font.name = CODE_FONT
            _rfonts(r._r, CODE_FONT)
            r.font.size = Pt(size)
            r.font.color.rgb = BLACK
        sp = self.doc.add_paragraph()
        sp.paragraph_format.first_line_indent = Inches(0)
        sp.paragraph_format.space_after = Pt(6)

    # ------------------------------------------------------------------ notas (A COLOR)
    def note(self, title, text, kind="info"):
        fills = {"info": ("E6F6FA", "06B6D4"), "warn": ("FEF3C7", "F59E0B"), "ok": ("DCFCE7", "16A34A")}
        fill, edge = fills[kind]
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_margins(t, 90, 90, 160, 160)
        tblPr = t._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        e = OxmlElement("w:left")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "30")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), edge)
        b.append(e)
        tblPr.append(b)
        c = t.rows[0].cells[0]
        c.width = Cm(16)
        cell_bg(c, fill)
        par = c.paragraphs[0]
        par.paragraph_format.first_line_indent = Inches(0)
        par.paragraph_format.space_after = Pt(3)
        r = par.add_run(title)
        r.bold = True
        r.font.name = HEAD_FONT
        _rfonts(r._r, HEAD_FONT)
        r.font.size = Pt(11)
        r.font.color.rgb = BLACK
        p2 = c.add_paragraph()
        p2.paragraph_format.first_line_indent = Inches(0)
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.2
        self._inline(p2, text, size=11, color=BLACK)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.first_line_indent = Inches(0)
        sp.paragraph_format.space_after = Pt(6)

    def page_break(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)
        p.add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------------ índices
    def toc_block(self, title, instr, placeholder, page_break=True):
        h = self.doc.add_paragraph()
        h.paragraph_format.page_break_before = page_break
        h.paragraph_format.first_line_indent = Inches(0)
        r = h.add_run(title)
        r.bold = True
        r.font.size = Pt(16)
        r.font.name = HEAD_FONT
        _rfonts(r._r, HEAD_FONT)
        r.font.color.rgb = BLACK
        para_border(h, "bottom", "000000", 12, 6)
        h.paragraph_format.space_after = Pt(14)
        par = self.doc.add_paragraph()
        par.paragraph_format.first_line_indent = Inches(0)
        add_field(par, instr, placeholder)

    # ------------------------------------------------------------------ secciones / encabezados / pie
    def _pgnum(self, section, fmt, start=None):
        sp = section._sectPr
        pg = sp.find(qn("w:pgNumType"))
        if pg is None:
            pg = OxmlElement("w:pgNumType")
            sp.append(pg)
        pg.set(qn("w:fmt"), fmt)
        if start is not None:
            pg.set(qn("w:start"), str(start))

    def _header_footer(self, section, running_head=True, header_text="LeoFit · Informe Final APF3"):
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        hp = section.header.paragraphs[0]
        hp.paragraph_format.first_line_indent = Inches(0)
        for r in list(hp.runs):
            r._r.getparent().remove(r._r)
        if running_head:
            r1 = hp.add_run(header_text + "\t")
            r1.font.size = Pt(9)
            r1.font.name = HEAD_FONT
            _rfonts(r1._r, HEAD_FONT)
            r1.font.color.rgb = DARK_GREY

            r2 = hp.add_run("UTP · Curso Integrador II: Software")
            r2.font.size = Pt(9)
            r2.font.name = HEAD_FONT
            _rfonts(r2._r, HEAD_FONT)
            r2.font.color.rgb = DARK_GREY

            para_border(hp, "bottom", "737373", 4, 4)
            hp.paragraph_format.tab_stops.add_tab_stop(Cm(16), alignment=2)

        fp = section.footer.paragraphs[0]
        fp.paragraph_format.first_line_indent = Inches(0)
        for r in list(fp.runs):
            r._r.getparent().remove(r._r)
        if running_head:
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_field(fp, "PAGE", "1", size=10, color=BLACK)
            fp.runs[-1].bold = True
            fp.runs[-1].font.name = BODY_FONT
            _rfonts(fp.runs[-1]._r, BODY_FONT)

    def new_section(self, kind="front"):
        """kind: front (romano) | body (arábigo, reinicia en 1)."""
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        if kind == "front":
            self._pgnum(sec, "lowerRoman", 1)
        else:
            self._pgnum(sec, "decimal", 1)
        self._header_footer(sec, running_head=(kind == "body"))
        if kind == "front":
            sec.header.is_linked_to_previous = False
            self._header_footer(sec, running_head=False)
            # pie con número romano
            fp = sec.footer.paragraphs[0]
            fp.paragraph_format.first_line_indent = Inches(0)
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_field(fp, "PAGE", "i", size=10, color=BLACK)
            fp.runs[-1].font.name = BODY_FONT
            _rfonts(fp.runs[-1]._r, BODY_FONT)
        return sec

    def finish_cover_section(self):
        s0 = self.doc.sections[0]
        self._header_footer(s0, running_head=False)

    def save(self, path):
        self.doc.save(path)
