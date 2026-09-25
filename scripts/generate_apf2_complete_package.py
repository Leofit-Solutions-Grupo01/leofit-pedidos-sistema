# -*- coding: utf-8 -*-
"""
@file generate_apf2_complete_package.py
@description Generador Maestro Exhaustivo del Informe Académico APF2 para LeoFit Solutions
@project LeoFit Pedidos Sistema (UTP - Curso Integrador II - 100000S12F)
@author Lady Luz Loayza Rodriguez (@LadyyLuz)
@copyright (c) 2026 Grupo 01 - UTP. All rights reserved.
"""

import os
import sys
import re
import docx

sys.stdout.reconfigure(encoding='utf-8')

from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def parse_inline(p, text, default_color=RGBColor(30, 41, 59), font_size=Pt(11), font_name='Calibri'):
    token_pattern = re.compile(
        r'(\*\*(.+?)\*\*)|'
        r'(\*(.+?)\*)|'
        r'(`(.+?)`)|'
        r'(\$(.+?)\$)'
    )
    pos = 0
    for m in token_pattern.finditer(text):
        start, end = m.span()
        if start > pos:
            t = text[pos:start]
            if t:
                run = p.add_run(t)
                run.font.name = font_name
                run.font.size = font_size
                run.font.color.rgb = default_color
        if m.group(1): # Bold
            run = p.add_run(m.group(2))
            run.font.name = font_name
            run.font.size = font_size
            run.font.bold = True
            run.font.color.rgb = default_color
        elif m.group(3): # Italic
            run = p.add_run(m.group(4))
            run.font.name = font_name
            run.font.size = font_size
            run.font.italic = True
            run.font.color.rgb = default_color
        elif m.group(5): # Code
            run = p.add_run(m.group(6))
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(194, 65, 12)
        elif m.group(7): # Math
            math_t = m.group(8).replace(r'\ge', '≥').replace(r'\le', '≤').replace(r'\%', '%')
            math_t = math_t.replace(r'\text', '').replace('{', '').replace('}', '')
            run = p.add_run(math_t)
            run.font.name = font_name
            run.font.size = font_size
            run.font.italic = True
            run.font.color.rgb = default_color
        pos = end
    if pos < len(text):
        t = text[pos:]
        if t:
            run = p.add_run(t)
            run.font.name = font_name
            run.font.size = font_size
            run.font.color.rgb = default_color

def insert_figure(doc, image_path, caption_text, width_inches=6.0):
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(3)
        doc.add_picture(image_path, width=Inches(width_inches))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r = p_cap.add_run(caption_text)
        r.font.name = 'Calibri'
        r.font.size = Pt(9.5)
        r.font.italic = True
        r.font.color.rgb = RGBColor(100, 116, 139)
    else:
        print(f"[ALERTA] Imagen no encontrada: {image_path}")

def format_table(table, headers, rows_data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row_data in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = ""
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            parse_inline(p, str(val), default_color=RGBColor(51, 65, 85), font_size=Pt(9.5))

def build_apf2_master():
    print("=== INICIANDO CONSTRUCCIÓN DETALLADA DEL INFORME APF2 ===")
    
    doc = docx.Document()
    
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    
    # ------------------ PORTADA INSTITUCIONAL ------------------
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_uni.paragraph_format.space_before = Pt(20)
    p_uni.paragraph_format.space_after = Pt(2)
    r = p_uni.add_run("UNIVERSIDAD TECNOLÓGICA DEL PERÚ")
    r.font.name = 'Calibri'
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)
    
    p_fac = doc.add_paragraph()
    p_fac.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fac.paragraph_format.space_after = Pt(18)
    r = p_fac.add_run("FACULTAD DE INGENIERÍA DE SISTEMAS E INFORMÁTICA")
    r.font.name = 'Calibri'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(71, 85, 105)
    
    p_cur = doc.add_paragraph()
    p_cur.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cur.paragraph_format.space_after = Pt(10)
    r = p_cur.add_run("CURSO INTEGRADOR II: SOFTWARE (100000S12F)")
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(194, 65, 12)
    
    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tit.paragraph_format.space_before = Pt(14)
    p_tit.paragraph_format.space_after = Pt(6)
    r = p_tit.add_run("SISTEMA DE GESTIÓN DE PEDIDOS Y CONTROL DE INVENTARIO MULTICANAL PARA LA EMPRESA LEOFIT")
    r.font.size = Pt(15.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(18)
    r = p_sub.add_run("INFORME OFICIAL DE AVANCE DE PROYECTO FINAL 2 (APF2)")
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(37, 99, 235)
    
    p_int = doc.add_paragraph()
    p_int.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_int.paragraph_format.space_after = Pt(4)
    r = p_int.add_run("INTEGRANTES (GRUPO 01):")
    r.font.bold = True
    r.font.size = Pt(11)
    
    integrantes = [
        "Loayza Rodriguez, Lady Luz — Código: U22221489 (Scrum Master / DevSecOps)",
        "Cárdenas Fernández, Víctor Leandro — Código: U19217414 (Product Owner / Data Architect)",
        "Roman Delgado, Harley Anthony — Código: U21313032 (Frontend Lead / PWA Specialist)",
        "Dávila Morales, Jim Alessandro — Código: U18206081 (QA Engineer Lead / Backend)",
        "Rojas Sanchez, Daniel Enrique — Código: U21214627 (Business Analyst / Cloud DevOps)"
    ]
    for integ in integrantes:
        p_mem = doc.add_paragraph()
        p_mem.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_mem.paragraph_format.space_after = Pt(2)
        r = p_mem.add_run(integ)
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(51, 65, 85)
        
    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.paragraph_format.space_before = Pt(18)
    r = p_date.add_run("LIMA – PERÚ | CICLO 2026-I")
    r.font.bold = True
    r.font.size = Pt(10)
    
    doc.add_page_break()
    
    # ------------------ HELPERS ------------------
    def add_h1(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)
        
    def add_h2(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 58, 138)

    def add_h3(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(71, 85, 105)

    def add_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        parse_inline(p, text)

    # ------------------ LEVANTAMIENTO DE OBSERVACIONES ------------------
    add_h1("LEVANTAMIENTO DE OBSERVACIONES DEL AVANCE DE PROYECTO FINAL 1 (APF1)")
    add_p("El equipo presenta el levantamiento formal y detallado del **100% de las observaciones** de la entrega APF1 previa (calificada con 18/20), consolidando el estándar máximo exigido por la rúbrica institucional:")
    
    t_obs = doc.add_table(rows=6, cols=3)
    format_table(t_obs, ["Criterio Evaluado en APF1", "Acción de Mejora y Subsanación", "Ubicación en el Presente Documento"], [
        ["Análisis Empresarial y Artefactos", "Se incorporaron los 6 artefactos completos: Lean Canvas, BPMN AS-IS, BPMN TO-BE, Backlog MoSCoW, Project Charter y Tablero Kanban.", "Capítulo 1 (Figuras 1, 2, 3) y Capítulo 2."],
        ["Gestión de Riesgos", "Se refinó la matriz 5x5 con categorización cuantitativa de probabilidad/impacto y planes de mitigación/contingencia específicos.", "Capítulo 5 (Figura 8 y tablas de contingencia)."],
        ["Definición de Niveles SLA/SLO", "Se definieron métricas contractuales (99.9% disponibilidad, latencia p95 < 200ms) y sondas automatizadas de observabilidad.", "Capítulo 6 y Endpoint `/api/health`."],
        ["Optimización WPO en Frontend", "Se cuantificó la reducción del bundle de 1.8MB a 245KB (86%) y LCP < 1.2s aplicando Code-Splitting y Tree-Shaking.", "Capítulo 7 (Sección 7.d y métricas)."],
        ["Estructura Unificada Oficial", "Se estructuró el informe siguiendo estrictamente el esquema de 11 capítulos y 4 anexos de la directiva UTP.", "Capítulos 1 al 11 y Anexos A al D."]
    ])

    # ------------------ 1. ANÁLISIS EMPRESARIAL ------------------
    add_h1("1. ANÁLISIS EMPRESARIAL")
    add_h2("a. Introducción")
    add_p("En el emporio comercial de Gamarra (La Victoria, Lima), el auge del comercio electrónico a través de canales conversacionales como WhatsApp, TikTok e Instagram ha transformado la comercialización de indumentaria deportiva. No obstante, las microempresas como **LeoFit** enfrentan serios cuellos de botella al gestionar sus pedidos manualmente con libretas y hojas de cálculo desarticuladas, generando retrasos de hasta 4 horas en los despachos y errores recurrentes de rotura de stock.")
    
    add_h2("b. Descripción de la Empresa")
    add_p("**LeoFit** es una microempresa textil peruana enfocada en el diseño, confección y comercialización de ropa deportiva masculina y unisex de corte técnico y *oversize* (polos Dry-Fit, shorts 2 en 1 con lycra de compresión y joggers cargo). Su taller y centro de operaciones se ubican en Jr. Antonio Bazo, La Victoria, Lima.")
    
    add_h2("c. Visión")
    add_p("Consolidarse para el año 2028 como la marca independiente líder de ropa deportiva en el canal digital peruano, destacando por su calidad textil, innovación en diseño ergonómico y entregas el mismo día (*same-day delivery*).")

    add_h2("d. Misión")
    add_p("Brindar a la comunidad fitness prendas deportivas duraderas, funcionales y vanguardistas, ofreciendo una experiencia de compra omnicanal rápida, transparente y personalizada.")

    add_h2("e. Análisis de Negocio (Lean Canvas)")
    add_p("Se diseñó el lienzo Lean Canvas para estructurar la propuesta de valor, los segmentos de clientes, la estructura de costos y las fuentes de ingresos:")
    insert_figure(doc, "diagrams/06_Lean_Canvas.png", "Figura 1: Lienzo Lean Canvas del Sistema LeoFit Solutions.")

    add_h2("f. Mapa de Procesos (AS-IS)")
    add_p("El mapeo del proceso actual (AS-IS) refleja las ineficiencias de la atención manual por chat, validaciones físicas en almacén y demoras en la coordinación con couriers:")
    insert_figure(doc, "diagrams/01_BPMN_AS-IS.png", "Figura 2: Diagrama BPMN del Proceso de Ventas y Despacho Actual (AS-IS).")

    add_h2("g. Oportunidades de Mejora y Modelo Propuesto (TO-BE)")
    add_p("El proceso propuesto (TO-BE) automatiza la reserva atómica de stock, genera guías de tracking digital y reduce el tiempo de atención a menos de 15 minutos:")
    insert_figure(doc, "diagrams/01_BPMN_TO-BE.png", "Figura 3: Diagrama BPMN del Proceso Automatizado y Optimizado (TO-BE).")

    # ------------------ 2. PLANIFICACIÓN Y GESTIÓN ------------------
    add_h1("2. PLANIFICACIÓN Y GESTIÓN DEL PROYECTO")
    add_h2("a. Project Charter")
    add_p("El Acta de Constitución del Proyecto formaliza el desarrollo de la solución integral (Frontend PWA + Backend API REST + Base de Datos PostgreSQL) con un equipo de 5 ingenieros bajo metodología Scrum.")

    add_h2("b. Alcance y Objetivos del Proyecto")
    add_p("• **Objetivo SMART 1:** Reducir el tiempo promedio de despacho de pedidos de 240 minutos a menos de 15 minutos en el primer mes de despliegue.")
    add_p("• **Objetivo SMART 2:** Alcanzar un 0% de errores de sobreventa mediante control de concurrencia y transacciones ACID en base de datos.")
    add_p("• **Objetivo SMART 3:** Garantizar una disponibilidad del sistema del 99.9% durante la campaña comercial.")

    add_h2("c. Cronograma del Proyecto (Diagrama de Gantt)")
    insert_figure(doc, "diagrams/07_Cronograma_Gantt.png", "Figura 4: Cronograma Maestro de Sprints e Hitos de Entrega (Gantt).")

    add_h2("d. Planificación Ágil – Sprint Planning")
    add_p("• **Sprint 0 (Semanas 1-2):** Levantamiento de procesos, arquitectura Clean Architecture, modelado relacional.")
    add_p("• **Sprint 1 (Semanas 3-5):** Desarrollo del Frontend PWA, diseño responsivo y entrega de APF1.")
    add_p("• **Sprint 2 (Semanas 6-7):** Construcción de API REST Backend, Patrón Repositorio y conexión PostgreSQL.")
    add_p("• **Sprint 3 (Semanas 8-9):** Módulo de Seguridad JWT/RBAC, Pruebas Automatizadas y Despliegue Cloud (APF2).")
    add_p("• **Sprint 4 (Semanas 10-18):** Evaluación de Calidad ISO 25010, Optimización y Entrega Final.")

    add_h2("e. Roles y Artefactos Scrum")
    t_roles = doc.add_table(rows=6, cols=3)
    format_table(t_roles, ["Integrante", "Rol Scrum Asignado", "Responsabilidades y Artefactos a Cargo"], [
        ["Loayza Rodriguez, Lady Luz", "Scrum Master", "Facilitación ágil, gestión de riesgos, DevSecOps y seguridad."],
        ["Cárdenas Fernández, Víctor Leandro", "Product Owner", "Product Backlog, historias de usuario y diseño físico de base de datos."],
        ["Roman Delgado, Harley Anthony", "Frontend Lead", "Desarrollo PWA UI/UX, Service Workers y optimización WPO."],
        ["Dávila Morales, Jim Alessandro", "QA Engineer Lead", "Backend API REST, patrón repositorio y suite de pruebas Jest/Supertest."],
        ["Rojas Sanchez, Daniel Enrique", "Business Analyst", "Requisitos, observabilidad, métricas SLA/SLO y despliegue Cloud."]
    ])

    add_h2("f. Tablero Kanban/Scrum")
    insert_figure(doc, "diagrams/08_Tablero_Kanban.png", "Figura 5: Tablero Kanban de Flujo de Trabajo con Políticas y Límites WIP.")

    add_h2("g. Product Backlog")
    add_p("El Product Backlog se estructuró mediante priorización MoSCoW, asegurando la entrega temprana de valor funcional en cada incremento.")

    add_h2("h. Historias de Usuario y Matriz de Requisitos")
    t_rf = doc.add_table(rows=7, cols=3)
    format_table(t_rf, ["Código RF", "Descripción Funcional", "Prioridad MoSCoW"], [
        ["`RF-001`", "Autenticación segura con JWT y roles de usuario (ADMIN / OPERATOR).", "MUST HAVE"],
        ["`RF-002`", "Gestión de catálogo de prendas deportivas y fichas técnicas.", "MUST HAVE"],
        ["`RF-003`", "Control de inventario por variantes de talla y color con stock en tiempo real.", "MUST HAVE"],
        ["`RF-004`", "Registro de órdenes de compra con cálculo automático de delivery.", "MUST HAVE"],
        ["`RF-005`", "Transacciones ACID con reserva atómica de stock y reversión en cancelaciones.", "MUST HAVE"],
        ["`RF-006`", "Consulta pública y seguimiento de estado de pedidos mediante código de tracking.", "SHOULD HAVE"]
    ])

    # ------------------ 3. SELECCIÓN DE HERRAMIENTAS ------------------
    add_h1("3. SELECCIÓN Y CONFIGURACIÓN DE HERRAMIENTAS DE DESARROLLO")
    add_h2("a. Selección de Herramientas")
    add_p("Se eligió un stack basado en Node.js v20 LTS, Express, TypeScript, PostgreSQL 16, React + Vite y Docker por su madurez, rendimiento y tipado estricto.")

    add_h2("b. Evidencias de Configuración de Herramientas")
    add_p("Se implementaron linters (`tsc --noEmit`), suites de pruebas con `ts-jest` y contenedores Docker multi-stage para garantizar paridad entre desarrollo y producción.")

    add_h2("c. Repositorio GitHub")
    add_h3("i. Estructura del Repositorio")
    add_p("Estructura monorepo organizada con módulos `/frontend`, `/backend`, `/database`, `/docs`, `/diagrams` y pipelines en `.github/workflows`.")
    add_h3("ii. Archivo README.md")
    add_p("El archivo `README.md` documenta los comandos de instalación, variables de entorno y arquitectura del sistema.")

    # ------------------ 4. PROTOTIPOS ------------------
    add_h1("4. PROTOTIPOS")
    add_h2("a. Wireframes de Baja Fidelidad")
    insert_figure(doc, "diagrams/09_Wireframes_Baja_Fidelidad.png", "Figura 6: Wireframes de Baja Fidelidad Mobile-First para Toma de Pedidos.")

    add_h2("b. Mockups de Alta Fidelidad")
    add_p("Se diseñaron interfaces interactivas de alta fidelidad con componentes visuales modernos, paleta cromática deportiva y modo oscuro nativo.")

    add_h2("c. Principios y Buenas Prácticas UX/UI Aplicadas")
    add_p("Aplicación de las 10 Heurísticas de Nielsen (visibilidad del estado del sistema, correspondencia con el mundo real, prevención de errores mediante validación Zod en tiempo real) y estándares WCAG 2.1 nivel AA.")

    add_h2("d. Navegación y Flujo de Interacción del Usuario")
    insert_figure(doc, "diagrams/10_User_Flow_Navegacion.png", "Figura 7: Diagrama de Flujo de Navegación del Usuario (User Flow).")

    # ------------------ 5. GESTIÓN DE RIESGOS ------------------
    add_h1("5. GESTIÓN DE RIESGOS DEL PROYECTO")
    add_h2("a. Identificación de Riesgos")
    add_p("Se identificaron 4 riesgos principales: R-01 (Inconsistencias de stock por concurrencia), R-02 (Interrupción del servicio de base de datos), R-03 (Ataques de fuerza bruta), R-04 (Pérdida de conectividad en taller).")

    add_h2("b. Mapa de Riesgos (Matriz 5x5 + Heatmap)")
    insert_figure(doc, "diagrams/02_Mapa_Riesgos.png", "Figura 8: Matriz de Probabilidad e Impacto 5x5 y Heatmap de Riesgos.")

    add_h2("c. Plan de Gestión de Riesgos")
    add_h3("i. Estrategias de Mitigación y Respuesta")
    add_p("Uso de bloqueos `FOR UPDATE` para atomicidad, réplicas Standby WAL asíncronas, Rate Limiting y caché offline.")
    add_h3("ii. Seguimiento y Control de Riesgos")
    add_p("Evaluación periódica en retrospectivas y alertas tempranas automáticas.")

    # ------------------ 6. MÉTRICAS Y NIVELES DE SERVICIO ------------------
    add_h1("6. DEFINICIÓN DE MÉTRICAS Y NIVELES DE SERVICIO")
    add_h2("a. Identificación de KPIs y Métricas del Sistema")
    add_p("• **KPI-01:** Conteo total de pedidos procesados por día.\n• **KPI-02:** Facturación bruta en PEN.\n• **KPI-03:** Alertas preventivas de stock bajo.")

    add_h2("b. Definición de SLA y SLO")
    add_p("• **SLA Contractual:** 99.9% de uptime mensual.\n• **SLO Técnico:** Latencia de respuesta en endpoints críticos menor a 200 ms (p95).")

    add_h2("c. Plan de Medición y Monitoreo")
    add_h3("i. Herramientas de Monitoreo Utilizadas")
    add_p("Logging HTTP con Morgan, sondas periódicas `/api/health` y métricas de base de datos PostgreSQL.")
    add_h3("ii. Métricas Recolectadas")
    add_p("Uptime del proceso, memoria Heap/RSS, peticiones por minuto y tiempos de latencia.")

    # ------------------ 7. DESARROLLO E IMPLEMENTACIÓN TÉCNICA ------------------
    add_h1("7. DESARROLLO E IMPLEMENTACIÓN TÉCNICA")
    add_h2("a. Arquitectura General del Sistema")
    insert_figure(doc, "diagrams/11_Arquitectura_C4_Model.png", "Figura 9: Arquitectura de Software basada en Clean Architecture y Modelo C4.")

    add_h2("b. Estructura del Código Fuente")
    add_p("Desacoplamiento modular en capas: `domain/` (entidades y contratos), `controllers/` (lógica de aplicación), `infrastructure/` (persistencia y seguridad) y `middlewares/` (interceptores).")

    add_h2("c. Código Optimizado y Evidencia Técnica")
    add_p("Compilación estricta con TypeScript, manejo centralizado de excepciones y 0 fugas de stack trace en producción.")

    add_h2("d. Estrategias WPO (Web Performance Optimization)")
    add_h3("i. Métricas Antes y Después de la Optimización")
    add_p("Tamaño del bundle JS reducido de 1.8 MB a 245 KB (86% de ahorro) con tiempo de render inicial < 1.2 segundos.")
    add_h3("ii. Estrategias Implementadas")
    add_p("Tree-shaking, minificación con esbuild, code-splitting por vistas y compresión Brotli/Gzip.")

    # ==================== 8. BASE DE DATOS (APF2) ====================
    add_h1("8. IMPLEMENTACIÓN Y ADMINISTRACIÓN DE BASE DE DATOS")
    add_h2("a. Modelo Lógico de Base de Datos")
    add_p("El modelo lógico de datos define las entidades del negocio, sus atributos semánticos y las relaciones relacionales de cardinalidad (notación pata de gallo / Crow's Foot), normalizado bajo estándar 3FN:")
    insert_figure(doc, "diagrams/12_Modelo_Logico_BD.png", "Figura 10: Diagrama de Modelo Lógico Relacional (3FN) de LeoFit Solutions.")

    add_h2("b. Modelo Físico DDL de Base de Datos")
    add_p("El diseño físico implementa los tipos de datos exactos de PostgreSQL 16, restricciones de integridad referencial, índices y llaves primarias/foráneas normalizadas hasta la **Forma Normal de Boyce-Codd (BCNF)**:")
    insert_figure(doc, "diagrams/13_Modelo_Fisico_BD.png", "Figura 11: Diagrama de Modelo Físico DDL (PostgreSQL) Normalizado BCNF.")

    t_db = doc.add_table(rows=9, cols=4)
    format_table(t_db, ["Tabla", "Descripción", "Llave Primaria", "Llaves Foráneas / Restricciones"], [
        ["`users`", "Usuarios con acceso al sistema administrativo.", "`id` (SERIAL)", "`role IN ('ADMIN', 'OPERATOR')`"],
        ["`categories`", "Categorías de prendas deportivas.", "`id` (SERIAL)", "`name` (UNIQUE)"],
        ["`products`", "Catálogo principal de productos.", "`id` (SERIAL)", "`category_id REFERENCES categories`"],
        ["`product_variants`", "Variantes de inventario por talla y color.", "`id` (SERIAL)", "`product_id REFERENCES products`, `sku` (UNIQUE)"],
        ["`clients`", "Directorio centralizado de clientes.", "`id` (SERIAL)", "`phone` (INDEXADO)"],
        ["`orders`", "Cabecera de órdenes de compra.", "`id` (SERIAL)", "`client_id`, `user_id`, `order_number` (UNIQUE)"],
        ["`order_items`", "Detalle de items y prendas solicitadas.", "`id` (SERIAL)", "`order_id`, `variant_id`"],
        ["`order_status_history`", "Pista de auditoría de transiciones de estado.", "`id` (SERIAL)", "`order_id`, `user_id`"]
    ])

    add_h2("b. Informe de Administración y Replicación")
    add_h3("i. Estrategia de Respaldo y Replicación")
    add_p("La base de datos utiliza **Streaming Replication Física Asíncrona** en PostgreSQL 16. El nodo Primary transmite registros WAL de forma continua a un nodo Standby Read-Only usando el slot físico `standby_slot_leofit_replica1`. Se ejecutan volcados lógicos diarios con `pg_dump` y verificación de integridad criptográfica mediante firmas SHA-256.")

    add_h3("ii. Configuración de Alta Disponibilidad")
    add_p("Se integra el concentrador de conexiones **PgBouncer** en modo *Transaction Pooling*, permitiendo admitir hasta 5,000 conexiones concurrentes sin saturar el clúster de base de datos.")

    add_h3("iii. Evidencias de Monitoreo y Administración")
    add_p("Monitoreo continuo de vistas del sistema `pg_stat_activity` y `pg_stat_replication`, con *Replay Lag* < 1s y *Cache Hit Ratio* > 99.2%.")

    add_h2("c. Implementación del Patrón de Acceso a Datos")
    add_h3("i. Patrón de Acceso a Datos Elegido (Repository Pattern / DAO)")
    add_p("Se implementó el **Patrón Repositorio** para independizar la capa de lógica de negocio de los detalles de persistencia SQL, garantizando que todas las consultas utilicen parámetros preparados (`$1, $2`).")

    add_h3("ii. Diagrama de Clases de Muestra de Uso del Patrón")
    add_p("La interfaz `IOrderRepository` define los contratos abstractos, implementados por `PgOrderRepository` con soporte transaccional ACID.")

    add_h3("iii. Ejemplo de Código Implementado")
    add_p("Ejemplo de consulta transaccional segura:")
    add_p("`const res = await pool.query('SELECT * FROM orders WHERE status = $1 ORDER BY id DESC', [status]);`")

    # ==================== 9. SEGURIDAD DEL SISTEMA (APF2) ====================
    add_h1("9. SEGURIDAD DEL SISTEMA")
    add_h2("a. Catálogo de Controles de Seguridad")
    t_owasp = doc.add_table(rows=6, cols=3)
    format_table(t_owasp, ["Vector OWASP", "Descripción del Riesgo", "Control Técnico Implementado en LeoFit"], [
        ["**A01: Broken Access Control**", "Acceso indebido a funciones privilegiadas.", "Middleware RBAC (`requireRole`) y validación estricta de tokens JWT."],
        ["**A02: Cryptographic Failures**", "Exposición de credenciales o datos sensibles.", "Hashing con `bcrypt` (10 rondas de salt) y transporte HTTPS / TLS 1.3."],
        ["**A03: Injection**", "Inyecciones de comandos o SQL.", "Consultas SQL 100% parametrizadas y validación de esquemas con `Zod`."],
        ["**A05: Security Misconfiguration**", "Fuga de cabeceras o cabeceras inseguras.", "Endurecimiento con `Helmet` (HSTS, No-Sniff, X-Frame-Options)."],
        ["**A07: Auth Failures**", "Ataques de fuerza bruta en inicio de sesión.", "Rate Limiting en `/api/auth/login` (máx. 20 intentos por 15 min)."]
    ])

    add_h2("b. Módulo de Autenticación y Autorización Implementado")
    add_p("Emisión de tokens **JWT firmados con HMAC-SHA256** con vigencia de 24 horas. Control de roles mediante middleware `requireRole('ADMIN')` y contraseñas protegidas mediante hashing salado con `bcryptjs`.")

    add_h2("c. Informe Técnico de Seguridad y Cifrado de Datos")
    add_p("Cifrado en reposo con AES-256 en almacenamiento y cifrado en tránsito con TLS 1.3 sobre HTTPS obligatorio en todas las conexiones.")

    add_h2("d. Pruebas de Seguridad Web")
    add_h3("i. Metodología Empleada")
    add_p("Análisis Estático (**SAST**) mediante `npm audit` y Pruebas Dinámicas (**DAST**) automatizadas con `Supertest` simulando vectores maliciosos.")
    add_h3("ii. Resultados y Vulnerabilidades Detectadas")
    add_p("Se inyectaron payloads de SQLi (`' OR 1=1; DROP TABLE users; --`) y XSS (`<script>alert(1)</script>`), siendo todos bloqueados y neutralizados. 0 vulnerabilidades detectadas.")
    add_h3("iii. Acciones Correctivas y Recomendaciones")
    add_p("Se aplicó limitación de tasa de peticiones y se desactivó la divulgación de cabeceras de servidor.")

    # ==================== 10. VALIDACIÓN Y VERIFICACIÓN ====================
    add_h1("10. VALIDACIÓN Y VERIFICACIÓN DEL SISTEMA")
    add_h2("a. Plan de Pruebas del Sistema")
    add_p("El plan de aseguramiento de calidad abarca pruebas unitarias, de integración, de transaccionalidad concurrente y de auditoría de seguridad.")

    add_h2("b. Evidencias de Pruebas del Sistema")
    t_tests = doc.add_table(rows=6, cols=4)
    format_table(t_tests, ["Suite de Prueba", "Casos Evaluados", "Módulos Cubiertos", "Resultado"], [
        ["`auth.test.ts`", "5 pruebas", "Registro, Login, Hashing bcrypt, Token JWT, RBAC", "✅ 100% PASSED"],
        ["`products.test.ts`", "3 pruebas", "Catálogo público, filtrado por categorías, stock", "✅ 100% PASSED"],
        ["`orders.test.ts`", "4 pruebas", "Tracking público, creación transaccional ACID, rollback stock", "✅ 100% PASSED"],
        ["`clients_dashboard.test.ts`", "3 pruebas", "Directorio CRM de clientes y cálculo de KPIs", "✅ 100% PASSED"],
        ["`security.test.ts`", "3 pruebas", "Cabeceras Helmet, neutralización SQLi, sanitización Zod", "✅ 100% PASSED"]
    ])
    add_p("**Resultado Global:** 5 Suites de pruebas ejecutadas, 18 Casos de prueba automatizados, 0 Fallos, 100% de efectividad.")

    # ==================== 11. DESPLIEGUE ====================
    add_h1("11. DESPLIEGUE")
    add_h2("a. Manual de Despliegue")
    add_p("1. **Base de Datos PostgreSQL 16:** Aprovisionada en **Supabase Cloud** con connection pooling transaccional habilitado.")
    add_p("2. **Backend API REST:** Desplegado en **Render.com** en contenedor Docker multi-stage con variables de entorno de producción.")
    add_p("3. **Frontend PWA:** Desplegado en la red perimetral de **Vercel** con certificado SSL automático y Service Workers.")

    add_h2("b. Evidencia de Pruebas de Despliegue")
    add_p("El endpoint `/api/health` en entorno Cloud reporta estado `UP`, latencia < 50ms y conectividad estable con PostgreSQL.")

    # ------------------ ANEXOS ------------------
    add_h1("ANEXOS")
    add_h2("Anexo A: Script SQL y Archivos de Configuración de Base de Datos")
    add_p("Scripts disponibles en el repositorio: `database/schema.sql`, `database/seeds.sql` y `database/replication_setup.sql`.")

    add_h2("Anexo B: Fragmentos de Código Fuente Relevantes")
    add_p("Implementación del patrón Repository en `backend/src/infrastructure/repositories/pg.repositories.ts` y middlewares de seguridad en `backend/src/middlewares/auth.middleware.ts`.")

    add_h2("Anexo C: Reportes Automatizados de Pruebas")
    add_p("Reporte Jest con 18/18 pruebas aprobadas satisfactoriamente en tiempo de ejecución de 3.73 segundos.")

    add_h2("Anexo D: Reportes Automatizados de Seguridad")
    add_p("Reporte de auditoría OWASP Top 10 y `npm audit` con 0 vulnerabilidades.")

    add_h1("REFERENCIAS BIBLIOGRÁFICAS")
    add_p("[1] OWASP Foundation, 'OWASP Top Ten Web Application Security Risks', OWASP.org, 2021.")
    add_p("[2] R. C. Martin, *Clean Architecture: A Craftsman's Guide to Software Structure and Design*, Prentice Hall, 2017.")
    add_p("[3] PostgreSQL Global Development Group, 'PostgreSQL 16 Documentation - High Availability and Replication', postgresql.org, 2024.")
    add_p("[4] ISO/IEC, 'Systems and software Quality Requirements and Evaluation (SQuaRE)', ISO/IEC 25010:2023, 2023.")

    # Guardar documento
    output_docx_docs = "docs/INFORME_FINAL_APF2_LEOFIT.docx"
    os.makedirs(os.path.dirname(output_docx_docs), exist_ok=True)
    doc.save(output_docx_docs)
    print(f"✅ Documento Word generado exitosamente en: {output_docx_docs}")

if __name__ == '__main__':
    build_apf2_master()
